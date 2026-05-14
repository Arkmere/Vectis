"""Focused checks for contextual ARCID/RM/airframe identifier classification."""

from __future__ import annotations

from pathlib import Path
import sys
from tempfile import TemporaryDirectory

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd
from openpyxl import load_workbook

from triage_engine import (
    CONFIG_REFERENCE_SPECS,
    classify_arcid_context,
    classify_callsign,
    derive_spanish_medical_registration,
    is_spanish_medical_registration_callsign,
    REQUIRED_COLUMNS,
    normalize_code,
    triage_file,
)


def _records(relative_path: str) -> list[dict[str, str]]:
    df = pd.read_csv(Path(relative_path), dtype=str, keep_default_na=False)
    for header in CONFIG_REFERENCE_SPECS[relative_path]:
        if header not in df.columns:
            df[header] = ""
    return [{column: str(row.get(column, "")).strip() for column in df.columns} for _, row in df.iterrows()]


CIVIL = _records("config/civil_registration_patterns.csv")
MILITARY = _records("config/military_serial_patterns.csv")
WORDLIKE = _records("config/wordlike_registrations.csv")
AIRCRAFT_CONTEXT = _records("config/aircraft_type_context.csv")
INTEREST_TYPES = set(pd.read_csv("config/interest_aircraft_types.csv", dtype=str, keep_default_na=False)["ATYP"].map(normalize_code)) - {""}


def ctx(arcid: str, rm: str = "", atyp: str = "") -> dict[str, object]:
    return classify_arcid_context(
        arcid,
        rm,
        atyp,
        civil_patterns=CIVIL,
        military_patterns=MILITARY,
        wordlike_registrations=WORDLIKE,
        aircraft_context=AIRCRAFT_CONTEXT,
        interest_types=INTEREST_TYPES,
    )


def _base_row(arcid: str, rm: str = "", atyp: str = "B738") -> dict[str, str]:
    row = {column: "" for column in REQUIRED_COLUMNS}
    row.update(
        {
            "ARCID": arcid,
            "ATYP": atyp,
            "RM": rm,
            "ADEP": "EGLL",
            "ADES": "EIDW",
            "IOBT": "2026-05-14T00:00:00Z",
        }
    )
    return row


def assert_workbook_unknown_operator_suppression() -> None:
    rows = [
        _base_row("QQQ1"),
        _base_row("HBJAZ", "HBJAZ"),
        _base_row("FEVER", "FEVER"),
        _base_row("MOOSE", "MOOSE", "SKRA"),
        _base_row("MEECCOF"),
        _base_row("MOOSE", "ONFILE", "C17"),
    ]
    with TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir)
        input_path = temp_path / "arcid_acceptance.csv"
        output_dir = temp_path / "output"
        pd.DataFrame(rows, columns=REQUIRED_COLUMNS).to_csv(input_path, index=False)
        result = triage_file(input_path, output_dir, ".")
        workbook = load_workbook(result.output_path, read_only=True)
        assert "23_ARCID_CLASS_DIAGNOSTICS" in workbook.sheetnames
        unknown_operator_sheet = workbook["05_UNKNOWN_OPERATOR_RANKED"]
        headers = [cell.value for cell in next(unknown_operator_sheet.iter_rows(max_row=1))]
        tricode_idx = headers.index("TRICODE")
        queued = {
            row[tricode_idx]
            for row in unknown_operator_sheet.iter_rows(min_row=2, values_only=True)
            if row[tricode_idx]
        }
        assert "QQQ" in queued
        assert not ({"HBJ", "FEV", "MOO", "MEE"} & queued)


def main() -> None:
    assert classify_callsign("BAW123")["CALLSIGN_FORM"] == "TRICODE_STYLE"
    assert classify_callsign("BAW123")["CALLSIGN_ROOT"] == "BAW"
    assert classify_callsign("RCH700")["CALLSIGN_FORM"] == "TRICODE_STYLE"
    assert classify_callsign("RCH700")["CALLSIGN_ROOT"] == "RCH"

    for value in ["HBJAZ", "FEVER", "MOOSE"]:
        assert classify_callsign(value)["CALLSIGN_FORM"] != "TRICODE_STYLE"
        assert ctx(value)["CALLSIGN_FORM"] != "TRICODE_STYLE"

    assert ctx("HBJAZ", "HBJAZ")["CALLSIGN_FORM"] == "AIRFRAME_IDENTIFIER_CALLSIGN"
    assert ctx("FEVER", "FEVER")["CALLSIGN_FORM"] == "AMBIGUOUS_REGISTRATION_OR_LONGFORM"
    assert ctx("MOOSE", "MOOSE")["CALLSIGN_FORM"] == "AMBIGUOUS_REGISTRATION_OR_LONGFORM"
    assert ctx("ZZ336", "ZZ336")["CALLSIGN_FORM"] == "MILITARY_SERIAL_CANDIDATE"

    med = ctx("MEECCOF")
    assert is_spanish_medical_registration_callsign("MEECCOF")
    assert derive_spanish_medical_registration("MEECCOF") == "EC-COF"
    assert med["CALLSIGN_FORM"] == "SPANISH_MEDICAL_REGISTRATION_CALLSIGN"
    assert med["ARCID_DERIVED_REGISTRATION"] == "EC-COF"
    assert med["CALLSIGN_ROOT"] != "MEE"

    assert ctx("BAW123", "ON")["RM_STATUS"] == "RM_ONFILE_SUPPRESSED"
    assert ctx("BAW123", "ONFILE")["RM_STATUS"] == "RM_ONFILE_SUPPRESSED"
    assert ctx("BAW123", "KNOWN")["RM_STATUS"] == "RM_KNOWN_SUPPRESSED"
    assert ctx("BAW123", "")["RM_STATUS"] == "RM_BLANK"

    assert ctx("MOOSE", "ONFILE", "C17")["CALLSIGN_FORM"] == "LONGFORM_NONSTANDARD"
    assert ctx("MOOSE", "MOOSE", "SKRA")["CALLSIGN_FORM"] == "AMBIGUOUS_REGISTRATION_OR_LONGFORM"

    assert_workbook_unknown_operator_suppression()

    print("arcid_context_validation_passed")


if __name__ == "__main__":
    main()
